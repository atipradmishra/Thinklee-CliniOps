import re
from datetime import datetime, date
from typing import List, Dict, Optional
from docx import Document
from app.utils.unstructured_ingest import normalize_text



def extract_events(filepath: str) -> List[Dict]:
    """
    Extract events directly from DOCX tables.
    Assumes table structure: Time | Description
    """

    try:
        doc = Document(filepath)
    except Exception as e:
        print(f"[extract_events] Failed to open DOCX: {e}")
        return []

    # -------------------------------------------------
    # 1️⃣ Extract date first (reuse your working function)
    # -------------------------------------------------
    file_date = extract_file_date_from_docx(doc)
    if not file_date:
        print("[extract_events] No date found → skipping event extraction")
        return []

    events: List[Dict] = []

    # -------------------------------------------------
    # 2️⃣ Loop through all tables
    # -------------------------------------------------
    for table in doc.tables:

        for row in table.rows:

            cells = [cell.text.strip() for cell in row.cells]

            if len(cells) < 2:
                continue

            time_candidate = cells[0]
            description = cells[1]

            # Skip header row
            if "tijd" in time_candidate.lower():
                continue

            # Detect time format (08:30 or 08.30)
            time_match = re.search(r"\b(\d{1,2}[:.]\d{2})\b", time_candidate)
            if not time_match:
                continue

            time_str = time_match.group(1).replace(".", ":")

            try:
                event_time = datetime.strptime(time_str, "%H:%M").time()
            except ValueError:
                continue

            event_datetime = datetime.combine(file_date, event_time)

            event_text = description.strip()
            if len(event_text) < 3:
                continue

            normalized_text = normalize_text(event_text)

            overlast_keywords = [
                "overlast",
                "geluidoverlast",
                "luidrucht",
                "baldadigheid",
                "hang jongeren",
                "dronken",
                "ruzie"
            ]

            contains_overlast = any(
                keyword in normalized_text for keyword in overlast_keywords
            )

            events.append({
                "event_date": file_date,
                "event_time": event_time,
                "event_datetime": event_datetime,
                "event_text": event_text,
                "normalized_text": normalized_text,
                "contains_overlast": contains_overlast
            })

    print(f"[extract_events] Found {len(events)} events")

    return events


# -------------------------------------------------
# Helper: Extract file date using DOCX structure
# -------------------------------------------------
def extract_file_date_from_docx(doc: Document) -> Optional[date]:
    """
    Robust DOCX date extraction.
    Checks:
    - Paragraphs
    - Tables
    - Headers
    - Dutch month names
    """

    dutch_months = {
        "januari": 1,
        "februari": 2,
        "maart": 3,
        "april": 4,
        "mei": 5,
        "juni": 6,
        "juli": 7,
        "augustus": 8,
        "september": 9,
        "oktober": 10,
        "november": 11,
        "december": 12
    }

    def try_parse_date(text: str) -> Optional[date]:
        text = text.strip()

        # Format: 12-07-2025
        match_numeric = re.search(r"\d{1,2}-\d{1,2}-\d{4}", text)
        if match_numeric:
            try:
                return datetime.strptime(match_numeric.group(), "%d-%m-%Y").date()
            except ValueError:
                pass

        # Format: 12 Juli 2025
        match_dutch = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", text)
        if match_dutch:
            day = int(match_dutch.group(1))
            month_name = match_dutch.group(2).lower()
            year = int(match_dutch.group(3))

            if month_name in dutch_months:
                return date(year, dutch_months[month_name], day)
        
        # Dutch month format
        match = re.search(r"(\d{1,2})[-\s]([a-zA-Z]+)[-\s](\d{4})", text)
        if match:
            day = int(match.group(1))
            month_name = match.group(2).lower()
            year = int(match.group(3))

            if month_name in dutch_months:
                return date(year, dutch_months[month_name], day)

        return None

    # -------------------------------------------------
    # 1️⃣ Check paragraphs
    # -------------------------------------------------
    for para in doc.paragraphs:
        parsed = try_parse_date(para.text)
        if parsed:
            return parsed

    # -------------------------------------------------
    # 2️⃣ Check tables
    # -------------------------------------------------
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parsed = try_parse_date(cell.text)
                if parsed:
                    return parsed

    # -------------------------------------------------
    # 3️⃣ Check headers
    # -------------------------------------------------
    for section in doc.sections:
        header = section.header
        for para in header.paragraphs:
            parsed = try_parse_date(para.text)
            if parsed:
                return parsed

    return None
