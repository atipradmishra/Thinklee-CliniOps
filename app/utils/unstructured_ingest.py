import os
import math
import csv
from pathlib import Path
import re
from typing import List
import unicodedata
from bs4 import BeautifulSoup
import fitz
from pptx import Presentation
from docx import Document
import pandas as pd
from lingua import LanguageDetectorBuilder

detector = LanguageDetectorBuilder.from_all_languages().build()

def detect_language(text: str):
    sample = text[:1000]
    language = detector.detect_language_of(sample)
    return language.iso_code_639_1.name.lower()


def normalize_text(text: str) -> str:
    if not isinstance(text, str):
        text = str(text)

    text = unicodedata.normalize("NFKC", text)
    text = text.lower()
    text = text.replace("\x00", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()

def extract_text_from_file(filepath: str) -> str:
    ext = Path(filepath).suffix.lower().lstrip(".")

    try:
        # ---------------- TXT / MD ----------------
        if ext in ("txt", "md"):
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                return normalize_text(f.read())

        # ---------------- HTML ----------------
        if ext in ("html", "htm"):
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                soup = BeautifulSoup(f.read(), "html.parser")
                for tag in soup(["script", "style", "noscript"]):
                    tag.extract()
                return normalize_text(soup.get_text(separator=" "))

        # ---------------- PDF ----------------
        if ext == "pdf":
            text = []
            try:
                with fitz.open(filepath) as doc:
                    for page in doc:
                        t = page.get_text("text")
                        if t.strip():
                            text.append(t)
            except Exception as e:
                print(f"[PDF] Primary read failed: {e}")

            return normalize_text("\n".join(text))

        # ---------------- DOCX ----------------
        if ext == "docx":
            doc = Document(filepath)
            blocks = []

            for p in doc.paragraphs:
                if p.text.strip():
                    blocks.append(p.text)

            for table in doc.tables:
                for row in table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text.strip()
                    ]
                    if cells:
                        blocks.append(" | ".join(cells))

            return normalize_text("\n".join(blocks))

        # ---------------- PPTX ----------------
        if ext == "pptx":
            prs = Presentation(filepath)
            slides_text = []

            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slides_text.append(shape.text)

            return normalize_text("\n".join(slides_text))

        # ---------------- CSV ----------------
        if ext == "csv":
            rows = []
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row in reader:
                    if any(cell.strip() for cell in row):
                        rows.append(" | ".join(row))
            return normalize_text("\n".join(rows))

        # ---------------- XLSX ----------------
        if ext in ("xlsx", "xls"):
            dfs = pd.read_excel(filepath, sheet_name=None)
            blocks = []

            for sheet, df in dfs.items():
                blocks.append(f"Sheet: {sheet}")
                blocks.append(df.astype(str).fillna("").to_string(index=False))

            return normalize_text("\n".join(blocks))

        # ---------------- FALLBACK ----------------
        with open(filepath, "rb") as f:
            raw = f.read()
        return normalize_text(raw.decode("utf-8", errors="ignore"))

    except Exception as e:
        print(f"[extract_text_from_file] Failed ({filepath}): {e}")
        return ""

    # finally:
    #     # Safe cleanup for temp uploads
    #     try:
    #         os.remove(filepath)
    #     except Exception:
    #         pass

def cosine(a: List[float], b: List[float]) -> float:
    if not a or not b or len(a) != len(b):
        return 0.0
    dot = sum(x*y for x, y in zip(a, b))
    na = math.sqrt(sum(x*x for x in a))
    nb = math.sqrt(sum(y*y for y in b))
    return dot / (na * nb + 1e-10)

def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 100) -> list[str]:
    if not text or not isinstance(text, str):
        return []

    chunks = []
    start = 0
    text_length = len(text)

    while start < text_length:
        end = min(start + chunk_size, text_length)
        chunk = text[start:end]
        chunks.append(chunk.strip())
        start += chunk_size - overlap

    return [c for c in chunks if c]

